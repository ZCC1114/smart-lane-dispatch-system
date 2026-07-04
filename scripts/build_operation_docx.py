from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "docs" / "操作文档.md"
OUTPUT_DOCX = ROOT / "docs" / "无锡硕放机场出租车蓄车池排队管理系统操作文档.docx"

ACCENT = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
BODY = RGBColor(31, 41, 55)
MUTED = RGBColor(100, 116, 139)
TABLE_HEADER = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
CALLOUT_BORDER = "BFD7EF"


def set_east_asia_font(run, name: str = "Microsoft YaHei") -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_paragraph_font(paragraph, size_pt: float = 11, color: RGBColor = BODY, bold: bool = False) -> None:
    for run in paragraph.runs:
        set_east_asia_font(run)
        run.font.size = Pt(size_pt)
        run.font.color.rgb = color
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_in: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), "9360")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    widths_dxa = [round(width * 1440) for width in widths_in]
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, color="CBD5E1", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def add_hyperlink_style_text(paragraph, text: str) -> None:
    pattern = re.compile(r"`([^`]+)`")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_east_asia_font(run)
        run = paragraph.add_run(match.group(1))
        set_east_asia_font(run)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)
        run.font.color.rgb = DARK_BLUE
        run.bold = True
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_east_asia_font(run)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.25

    h1 = styles["Heading 1"]
    h1.font.size = Pt(16)
    h1.font.color.rgb = ACCENT
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)

    h2 = styles["Heading 2"]
    h2.font.size = Pt(13)
    h2.font.color.rgb = ACCENT
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(7)

    h3 = styles["Heading 3"]
    h3.font.size = Pt(12)
    h3.font.color.rgb = DARK_BLUE
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(5)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "无锡硕放机场出租车蓄车池排队管理系统 | 操作文档"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_font(header, 9, MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("第 ")
    set_east_asia_font(run)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    footer._p.append(fld_begin)
    footer._p.append(instr)
    footer._p.append(fld_end)
    run = footer.add_run(" 页")
    set_east_asia_font(run)
    set_paragraph_font(footer, 9, MUTED)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(86)
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("无锡硕放机场出租车蓄车池排队管理系统")
    set_east_asia_font(r)
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(15, 23, 42)
    r.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("操作文档")
    set_east_asia_font(r)
    r.font.size = Pt(22)
    r.font.color.rgb = ACCENT
    r.bold = True

    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [1.35, 5.15])
    set_table_borders(table, "D6E3F0", "6")
    rows = [
        ("适用对象", "现场调度人员、值班管理员、运维人员"),
        ("文档内容", "登录与退出、菜单功能、图文操作说明、出入口红绿灯切换逻辑、常见操作流程"),
        ("截图来源", "系统本地演示环境，现场数据以实际运行界面为准"),
        ("维护位置", "docs/操作文档.md 与 docs/assets/operation-guide/"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], TABLE_HEADER)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                set_paragraph_font(paragraph, 10.5, BODY, bold=(cell is row.cells[0]))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(8)
    p.style = doc.styles["Normal"]
    r = p.add_run("阅读提示：")
    set_east_asia_font(r)
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    r.font.size = Pt(11)
    r2 = p.add_run("本文已将系统截图嵌入对应功能章节，便于现场对照界面操作。")
    set_east_asia_font(r2)
    r2.font.size = Pt(11)
    r2.font.color.rgb = BODY


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        parts = [part.strip() for part in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", part or "") for part in parts):
            continue
        rows.append(parts)
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    if cols == 3:
        widths = [1.1, 2.55, 2.85]
    elif cols == 2:
        widths = [1.7, 4.8]
    else:
        widths = [6.5 / cols] * cols
    set_table_geometry(table, widths)
    set_table_borders(table)
    for r_idx, row_data in enumerate(rows):
        for c_idx, text in enumerate(row_data):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            add_hyperlink_style_text(p, text)
            if r_idx == 0:
                set_cell_shading(cell, TABLE_HEADER)
                set_paragraph_font(p, 10.5, RGBColor(15, 23, 42), True)
            else:
                set_paragraph_font(p, 10.2, BODY)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)


def add_image(doc: Document, md_path: str, alt: str, figure_no: int) -> int:
    image_path = (SOURCE_MD.parent / md_path).resolve()
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(3)
    r = caption.add_run(f"图 {figure_no}  {alt}")
    set_east_asia_font(r)
    r.font.size = Pt(9.5)
    r.font.color.rgb = MUTED
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(str(image_path), width=Inches(6.3))
    return figure_no + 1


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    text = "\n".join(line.rstrip() for line in code_lines).strip()
    if not text:
        return
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    set_table_borders(table, "D8E2EE", "4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    run.font.color.rgb = DARK_BLUE
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    set_table_borders(table, CALLOUT_BORDER, "6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    add_hyperlink_style_text(p, text)
    set_paragraph_font(p, 10.5, DARK_BLUE)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)


def add_text_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    add_hyperlink_style_text(p, text)
    set_paragraph_font(p)


def build_docx() -> None:
    md = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    figure_no = 1
    i = 0
    in_code = False
    code_lines: list[str] = []
    while i < len(md):
        raw = md[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(md) and md[i].strip().startswith("|"):
                table_lines.append(md[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            figure_no = add_image(doc, image_match.group(2), image_match.group(1), figure_no)
            i += 1
            continue

        if stripped.startswith(">"):
            callout_text = stripped.lstrip(">").strip()
            if not callout_text.startswith("本文用于现场调度人员"):
                add_callout(doc, callout_text)
            i += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            if level == 1:
                i += 1
                continue
            style = "Heading 1" if level == 2 else "Heading 2"
            p = doc.add_paragraph(style=style)
            add_hyperlink_style_text(p, text)
            set_paragraph_font(p, 16 if level == 2 else 13, ACCENT, True)
            i += 1
            continue

        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        if bullet_match:
            p = doc.add_paragraph(style="List Bullet")
            add_hyperlink_style_text(p, bullet_match.group(1))
            set_paragraph_font(p)
            i += 1
            continue

        number_match = re.match(r"^\d+\.\s+(.+)$", stripped)
        if number_match:
            p = doc.add_paragraph(style="List Number")
            add_hyperlink_style_text(p, number_match.group(1))
            set_paragraph_font(p)
            i += 1
            continue

        add_text_paragraph(doc, stripped)
        i += 1

    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_docx()
